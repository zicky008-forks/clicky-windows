"""Build the monthly report from two exported trial-balance files, without Excel running.

    python -m engine.cli TB_T10.xlsx TB_T09.xlsx -o BaoCao_T10.xlsx

Each input's first sheet (or --sheet) is copied into the output next to the report,
so every formula and source link resolves inside one workbook. Needs openpyxl.
"""
import argparse
import csv
from decimal import Decimal
from pathlib import Path

from engine.build import build

SOURCE_CUR, SOURCE_PREV = "Goc ky nay", "Goc ky truoc"
NUMBER_FORMAT = '#,##0;(#,##0);"-"'


def read_rows(path: Path, sheet=None) -> list:
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as f:
            return [row for row in csv.reader(f)]
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def write_workbook(path: Path, sheets: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    for name, grid in sheets.items():
        ws = wb.create_sheet(name[:31])
        for row in grid:
            ws.append([float(v) if isinstance(v, Decimal) else v for v in row])
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, (int, float)) or (isinstance(cell.value, str) and cell.value.startswith(("=SUMIFS", "=C", "=D"))):
                    cell.number_format = NUMBER_FORMAT
    report = wb.worksheets[0]
    report["A1"].font = Font(bold=True, size=14)
    for cell in report[4]:
        cell.font = Font(bold=True)
    for col, width in zip("ABCDEFGHI", (44, 12, 18, 18, 18, 9, 14, 48, 16)):
        report.column_dimensions[col].width = width
    for row in report.iter_rows(min_row=5, min_col=6, max_col=6):
        row[0].number_format = "0.0%"
    wb.save(path)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("current")
    ap.add_argument("previous")
    ap.add_argument("-o", "--output", default="BaoCao_thang.xlsx")
    ap.add_argument("--sheet", help="sheet name in both inputs (default: first sheet)")
    args = ap.parse_args(argv)
    cur_rows = read_rows(Path(args.current), args.sheet)
    prev_rows = read_rows(Path(args.previous), args.sheet)
    cur_name = Path(args.current).stem[:31]
    prev_name = Path(args.previous).stem[:31]
    if cur_name == prev_name:
        cur_name, prev_name = SOURCE_CUR, SOURCE_PREV
    report = build(cur_rows, prev_rows, cur_name, prev_name)
    sheets = dict(report.sheets)
    sheets[cur_name] = cur_rows
    sheets[prev_name] = prev_rows
    write_workbook(Path(args.output), sheets)
    print(report.summary)
    print(f"Đã ghi {args.output} ({len(report.issues)} điểm kiểm tra)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
