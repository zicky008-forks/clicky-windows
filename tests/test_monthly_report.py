from decimal import Decimal

import pytest

from engine.build import build
from engine.checks import check
from engine.monthly_report import compare, compute, fmt
from engine.trial_balance import fold, parse_amount, read_trial_balance
from tb_samples import M, t9_rows, t10_rows


def values(figures):
    return {f.key: f.value / M for f in figures}


def test_amount_formats():
    assert parse_amount("1.234.567") == (Decimal(1234567), True)
    assert parse_amount("1,234,567.50") == (Decimal("1234567.50"), True)
    assert parse_amount("(2.000)") == (Decimal(-2000), True)
    assert parse_amount(1500.0) == (Decimal("1500.0"), False)
    assert parse_amount(None) == (Decimal(0), False)
    with pytest.raises(ValueError):
        parse_amount("abc")
    assert fold("Số dư đầu kỳ") == "so du dau ky"


def test_reads_misa_layout_and_marks_parents():
    tb = read_trial_balance(t10_rows(), "T10")
    assert tb.header_row == 5 and set(tb.columns) == {"open_dr", "open_cr", "dr", "cr", "close_dr", "close_cr"}
    by = {l.account: l for l in tb.lines}
    assert not by["511"].leaf and by["5111"].leaf and by["911"].leaf
    assert by["5113"].cr == 450 * M and by["5113"].row == tb.lines[0].row + list(by).index("5113")
    assert "Tổng cộng" not in [l.account for l in tb.lines]


def test_pnl_counts_children_once():
    v = values(compute(read_trial_balance(t10_rows())))
    assert (v["rev"], v["ded"], v["net_rev"], v["cogs"], v["gross"]) == (8850, 210, 8640, 5800, 2840)
    assert (v["selling"], v["admin"], v["op"], v["other"], v["pbt"], v["tax"], v["pat"]) == (1380, 690, 715, -20, 695, 139, 556)


def test_variance_names_the_accounts_that_moved():
    cur = compute(read_trial_balance(t10_rows()))
    prev = compute(read_trial_balance(t9_rows()))
    v = {x.key: x for x in compare(cur, prev)}
    assert v["pat"].change / M == 556 - 924
    assert v["selling"].change / M == 230
    assert v["selling"].drivers[0][0] == "6417" and v["selling"].drivers[0][2] / M == 120
    assert round(v["pat"].pct, 1) == Decimal("-39.8")


def test_clean_books_raise_no_issues():
    assert check(read_trial_balance(t10_rows())) == []
    assert check(read_trial_balance(t9_rows())) == []


def test_returns_booked_into_revenue_are_caught():
    issues = check(read_trial_balance(t10_rows(returns_into_511=37)))
    assert [i.code for i in issues] == ["closing_mismatch"]


def test_broken_trial_balance_is_flagged():
    rows = t10_rows()
    i = next(n for n, r in enumerate(rows) if r and r[0] == "6417")
    rows[i][4] += 1_000_000                       # one extra million on Phát sinh Nợ 6417
    rows.append(list(rows[i]))                    # and the account pasted twice
    codes = {x.code for x in check(read_trial_balance(rows))}
    assert {"unbalanced", "rollforward", "parent_sum", "duplicate"} <= codes


def test_text_amounts_are_reported():
    rows = t10_rows()
    i = next(n for n, r in enumerate(rows) if r and r[0] == "6418")
    rows[i][4] = fmt(Decimal(rows[i][4]))       # "200.000.000" typed as text
    tb = read_trial_balance(rows)
    assert {l.account: l for l in tb.lines}["6418"].dr == 200 * M
    assert [x.code for x in check(tb)] == ["text_amount"]


def test_report_sheets_use_traceable_formulas():
    report = build(t10_rows(), t9_rows(), "Goc T10", "Goc T09")
    grid = report.sheets["BC thang"]
    header = grid[3]
    assert header[:3] == ["Chỉ tiêu", "Tài khoản", "Kỳ này"]
    selling = next(r for r in grid if r and r[0] == "Chi phí bán hàng")
    assert selling[2] == "=SUMIFS('TB ky nay'!$E:$E,'TB ky nay'!$A:$A,\"641*\")"
    assert selling[7].startswith("=IF(ABS(C") and "1380000000" in selling[7]
    net = next(r for r in grid if r and r[0] == "Doanh thu thuần")
    assert net[2] == "=C5-C6"
    tb_sheet = report.sheets["TB ky nay"]
    assert [r[0] for r in tb_sheet[1:]].count("511") == 0 and tb_sheet[1][-1].startswith('=HYPERLINK("#\'Goc T10\'!A')
    assert report.summary.startswith("Lợi nhuận sau thuế kỳ này 556 triệu, giảm 368 triệu (-39,8%)")
    assert "giá vốn hàng bán tăng 500 triệu" in report.summary and "TK 6417" in report.summary


def test_monthly_report_skill_trigger():
    from office.report_skill import TRIGGER, pick_periods
    import re
    pattern = re.compile(TRIGGER, re.IGNORECASE)
    assert pattern.search("lập báo cáo quản trị tháng 10 giúp anh")
    assert pattern.search("lam bao cao thang")
    assert not pattern.search("giải thích báo cáo tháng này")
    found = [("T08", []), ("T09", []), ("T10", [])]
    assert [n for n, _ in pick_periods(found, "T09")] == ["T09", "T08"]
    assert [n for n, _ in pick_periods(found, "Ghi chú")] == ["T10", "T09"]


class FakeWorkbookBridge:
    def __init__(self, sheets, active):
        self.sheets, self.active, self.written = dict(sheets), active, {}

    def available(self):
        return True

    def sheet_names(self):
        return list(self.sheets)

    def sheet_rows(self, name):
        return self.sheets[name]

    def context(self):
        return {"sheet": self.active}

    def add_sheet(self, name):
        final = name if name not in self.sheets else name + " 2"
        self.sheets[final] = []
        return final

    def write(self, ref, grid):
        self.written[ref] = grid
        return True

    def format_numbers(self, ref, number_format=None):
        pass

    def activate(self, name):
        self.active = name


def test_skill_builds_report_inside_open_workbook():
    from office.report_skill import run
    bridge = FakeWorkbookBridge({"T09": t9_rows(), "T10": t10_rows(), "Ghi chú": [["memo"]], "BC thang": []}, "T10")
    summary = run(bridge)
    assert summary.startswith("Lợi nhuận sau thuế kỳ này 556 triệu")
    assert bridge.active == "BC thang 2"                      # existing sheet left alone
    report = bridge.written["'BC thang 2'!A1"]
    assert "'TB ky nay'" in next(r for r in report if r and r[0] == "Giá vốn hàng bán")[2]
    assert "'T09'!A" in bridge.written["'TB ky truoc'!A1"][1][-1]
    assert run(FakeWorkbookBridge({"T10": t10_rows()}, "T10")).startswith("Anh mở file Excel")


def test_cli_writes_one_self_contained_workbook(tmp_path):
    from openpyxl import load_workbook
    from engine.cli import main, write_workbook
    write_workbook(tmp_path / "TB_T10.xlsx", {"CDTK": t10_rows()})
    write_workbook(tmp_path / "TB_T09.xlsx", {"CDTK": t9_rows()})
    out = tmp_path / "BaoCao.xlsx"
    assert main([str(tmp_path / "TB_T10.xlsx"), str(tmp_path / "TB_T09.xlsx"), "-o", str(out)]) == 0
    wb = load_workbook(out)
    assert wb.sheetnames == ["BC thang", "Kiem tra", "TB ky nay", "TB ky truoc", "TB_T10", "TB_T09"]
    assert wb["BC thang"]["C5"].value.startswith("=SUMIFS('TB ky nay'!$F:$F")
